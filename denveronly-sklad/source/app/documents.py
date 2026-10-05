"""Генерація акта у PDF (reportlab) та DOCX (python-docx).

Аркуш 1 — акт: оренда + один рядок «Компенсація комунальних послуг» на загальну суму.
Аркуш 2 — «Деталізація комунальних послуг»: кожна послуга з показниками лічильників.
"""
import io
import os

from calc import MONTHS_UK_GEN, act_layout, amount_in_words, fmt_money, fmt_num

FONT_DIRS = ["/usr/share/fonts/truetype/dejavu", os.path.join(os.path.dirname(__file__), "fonts")]

MAIN_HEAD = ["№", "Найменування", "Од.", "К-сть", "Ціна, грн", "Сума, грн"]
DETAIL_HEAD = ["№", "Послуга", "Склад / лічильник", "Показники", "Од.", "К-сть", "Ціна, грн", "Сума, грн"]


def _font(name):
    for d in FONT_DIRS:
        p = os.path.join(d, name)
        if os.path.exists(p):
            return p
    raise FileNotFoundError(f"Шрифт {name} не знайдено (встановіть fonts-dejavu-core)")


def _act_date(act):
    """Дата акта зі знімка; для старих актів — останній день місяця."""
    if act["data"].get("act_date"):
        return act["data"]["act_date"]
    import calendar
    y, m = map(int, act["period"].split("-"))
    return f"{calendar.monthrange(y, m)[1]} {MONTHS_UK_GEN[m - 1]} {y} р."


def _ulabel(act):
    return act["data"].get("util_period_label") or act["data"]["period_label"]


def _party(p):
    rows = [p.get("name") or ""]
    if p.get("edrpou"):
        rows.append(f"ЄДРПОУ/РНОКПП: {p['edrpou']}")
    if p.get("address"):
        rows.append(p["address"])
    if p.get("iban"):
        rows.append(f"IBAN: {p['iban']}")
    if p.get("bank"):
        rows.append(p["bank"])
    return rows


def _company(act):
    c = act["data"]["company"]
    return {"name": c.get("company_name"), "edrpou": c.get("company_edrpou"),
            "address": c.get("company_address"), "iban": c.get("company_iban"),
            "bank": c.get("company_bank"), "signer": c.get("company_signer"),
            "position": c.get("company_signer_position"), "basis": c.get("company_basis")}


def _tenant(act):
    t = act["data"]["tenant"]
    return {"name": t.get("name"), "edrpou": t.get("edrpou"), "address": t.get("address"),
            "iban": t.get("iban"), "signer": t.get("contact"),
            "position": t.get("director_position"), "basis": t.get("basis")}


def short_name(full):
    """«Іваненко Іван Іванович» -> «Іваненко І.І.»; інше — без змін."""
    parts = (full or "").split()
    if len(parts) in (2, 3) and all(p[:1].isupper() for p in parts) and "." not in full:
        return parts[0] + " " + "".join(p[0] + "." for p in parts[1:])
    return full or ""


def _who(p, role):
    """«ТОВ «Альфа» (Орендар), від імені якого діє Директор Іваненко Іван Іванович на підставі Статуту»."""
    s = f"{p['name'] or ''} (далі — {role})"
    if p.get("signer"):
        s += f", від імені якого діє {p.get('position') or 'представник'} {p['signer']}"
        if p.get("basis"):
            s += f" на підставі {p['basis']}"
    return s


def _sign_line(p):
    pos = (p.get("position") or "").strip()
    return (pos + " " if pos else "") + "_______________ " + short_name(p.get("signer"))


def _contract(act):
    t = act["data"]["tenant"]
    if not t.get("contract_no"):
        return ""
    s = f"Договором № {t['contract_no']}"
    if t.get("contract_date"):
        s += f" від {t['contract_date']}"
    return s


def _intro(act):
    d = act["data"]
    contract = _contract(act)
    lay = act_layout(d)
    has_rent = any(l.get("group") == "rent" for l in d["lines"])
    parts = []
    if has_rent:
        parts.append(f"послуги з оренди за {d['period_label']}")
    if lay["util"]:
        parts.append(f"компенсацію комунальних послуг за {_ulabel(act)}")
    what = " та ".join(parts) if parts else f"послуги за {d['period_label']}"
    return (f"{_who(_company(act), 'Орендодавець')}, з одного боку, та {_who(_tenant(act), 'Орендар')}, "
            f"з другого боку, склали цей акт про те, що Орендодавцем надано, а Орендарем прийнято {what}"
            + (f" згідно з {contract}" if contract else "") + ":")


def _detail_intro(act):
    d = act["data"]
    contract = _contract(act)
    return (f"Орендодавець: {_company(act)['name'] or ''}. Орендар: {d['tenant']['name']}. "
            f"Період споживання: {_ulabel(act)}." + (f" Підстава: {contract}." if contract else ""))


def _totals_rows(act):
    tt = act["data"]["totals"]
    if tt.get("vat_rate"):
        return [("Разом без ПДВ:", fmt_money(tt["net"])), (f"ПДВ {tt['vat_rate']}%:", fmt_money(tt["vat"])),
                ("Всього з ПДВ:", fmt_money(tt["total"]))]
    return [("Всього (без ПДВ):", fmt_money(tt["total"]))]


def _detail_totals_rows(act, lay):
    vr = act["data"]["totals"].get("vat_rate")
    if vr:
        return [("Разом без ПДВ:", fmt_money(lay["util_sum"])), (f"ПДВ {vr}%:", fmt_money(lay["util_vat"])),
                ("Всього з ПДВ:", fmt_money(lay["util_total"]))]
    return [("Всього (без ПДВ):", fmt_money(lay["util_sum"]))]


def _words(act):
    tt = act["data"]["totals"]
    s = f"Загальна вартість: {amount_in_words(tt['total'])}"
    return s + (f", у т.ч. ПДВ {fmt_money(tt['vat'])} грн." if tt.get("vat_rate") else ", без ПДВ.")


def _detail_rows(lay):
    """Рядки таблиці деталізації (текстом)."""
    rows = []
    for i, l in enumerate(lay["util"], 1):
        if "service" in l:                                    # структуровані дані
            where = l.get("wh") or "—"
            if l.get("serial"):
                where += f", №{l['serial']}"
            if "cur" in l:
                readings = f"{fmt_num(l['prev'])} → {fmt_num(l['cur'])}"
                if (l.get("coef") or 1) != 1:
                    readings += f" ×{fmt_num(l['coef'])}"
            else:
                readings = "—"
            name = l["service"]
        else:                                                 # акти, сформовані до цієї версії
            name, where, readings = l["name"], "—", "—"
        rows.append([str(i), name, where, readings, l["unit"], fmt_num(l["qty"]),
                     fmt_money(l["price"]), fmt_money(l["sum"])])
    return rows


# ---------------- PDF ----------------

def act_pdf(act):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", _font("DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", _font("DejaVuSans-Bold.ttf")))

    base = ParagraphStyle("b", fontName="DejaVu", fontSize=9, leading=12)
    small = ParagraphStyle("s", parent=base, fontSize=8, leading=10)
    title = ParagraphStyle("t", parent=base, fontName="DejaVu-Bold", fontSize=13, leading=16, alignment=TA_CENTER)
    center = ParagraphStyle("c", parent=base, alignment=TA_CENTER)
    right = ParagraphStyle("r", parent=small, alignment=TA_RIGHT)
    bold = ParagraphStyle("bo", parent=base, fontName="DejaVu-Bold")
    bold_s = ParagraphStyle("bos", parent=small, fontName="DejaVu-Bold")

    d = act["data"]
    lay = act_layout(d)
    c, t = _company(act), _tenant(act)

    def header(lines):
        out = [Paragraph(lines[0], title)] + [Paragraph(x, center) for x in lines[1:]] + [Spacer(1, 3 * mm)]
        out.append(Table([[Paragraph(d.get("city") or "", base), Paragraph(_act_date(act), right)]],
                         colWidths=[90 * mm, 90 * mm]))
        return out + [Spacer(1, 3 * mm)]

    def table(head, rows, widths, num_from):
        data = [[Paragraph(h, bold_s) for h in head]]
        for r in rows:
            data.append([Paragraph(str(x), small) if j in (1, 2) else x for j, x in enumerate(r)])
        tb = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
        tb.setStyle(TableStyle([
            ("FONT", (0, 0), (-1, -1), "DejaVu", 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (num_from, 1), (-1, -1), "RIGHT"),
            ("ALIGN", (0, 1), (0, -1), "CENTER"),
        ]))
        return tb

    def totals(rows):
        tb = Table([[Paragraph(a, bold), Paragraph(b, bold)] for a, b in rows], colWidths=[155 * mm, 25 * mm])
        tb.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "RIGHT")]))
        return tb

    def signatures(full=True):
        if full:
            left = [Paragraph("Від Орендодавця:", bold)] + [Paragraph(x, small) for x in _party(c)]
            rightc = [Paragraph("Від Орендаря:", bold)] + [Paragraph(x, small) for x in _party(t)]
        else:
            left = [Paragraph("Від Орендодавця:", bold)]
            rightc = [Paragraph("Від Орендаря:", bold)]
        rows = [[left, rightc], ["", ""],
                [Paragraph(_sign_line(c), base), Paragraph(_sign_line(t), base)],
                [Paragraph("М.П.", small), Paragraph("М.П.", small)]]
        sg = Table(rows, colWidths=[90 * mm, 90 * mm], rowHeights=[None, 10 * mm, None, None])
        sg.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return sg

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm, title=f"Акт {act['number']}")

    # ---- аркуш 1: акт ----
    el = header([f"АКТ № {act['number']}", "приймання-передачі наданих послуг"])
    el += [Paragraph(_intro(act), base), Spacer(1, 3 * mm)]
    main_rows = [[str(i), l["name"], l["unit"], fmt_num(l["qty"]), fmt_money(l["price"]), fmt_money(l["sum"])]
                 for i, l in enumerate(lay["main"], 1)]
    el.append(table(MAIN_HEAD, main_rows, [8, 95, 17, 18, 21, 21], 3))
    el += [Spacer(1, 2 * mm), totals(_totals_rows(act)), Spacer(1, 3 * mm),
           Paragraph(_words(act), base), Spacer(1, 2 * mm)]
    if lay["util"]:
        el.append(Paragraph("Розрахунок суми компенсації комунальних послуг наведено в Деталізації "
                            "(додаток до цього акта).", base))
    el += [Paragraph(f"Форма оплати: {d['payment_label']}. Сторони претензій одна до одної не мають.", base),
           Spacer(1, 10 * mm), signatures()]

    # ---- аркуш 2: деталізація ----
    if lay["util"]:
        el.append(PageBreak())
        el += header(["ДЕТАЛІЗАЦІЯ КОМУНАЛЬНИХ ПОСЛУГ",
                      f"за {_ulabel(act)} — додаток до акта № {act['number']} від {_act_date(act)}"])
        el += [Paragraph(_detail_intro(act), base), Spacer(1, 3 * mm)]
        el.append(table(DETAIL_HEAD, _detail_rows(lay), [8, 45, 30, 27, 15, 17, 19, 19], 4))
        el += [Spacer(1, 2 * mm), totals(_detail_totals_rows(act, lay)), Spacer(1, 3 * mm),
               Paragraph(f"Сума компенсації комунальних послуг: {amount_in_words(lay['util_total'])}"
                         + (f", у т.ч. ПДВ {fmt_money(lay['util_vat'])} грн." if lay["util_vat"] else ", без ПДВ."),
                         base),
               Paragraph("Ціни вказано без ПДВ. Кількість за лічильниками = (поточні − попередні показники) "
                         "× коефіцієнт.", small),
               Spacer(1, 10 * mm), signatures(full=False)]

    doc.build(el)
    return buf.getvalue()


# ---------------- DOCX ----------------

def act_docx(act):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    d = act["data"]
    lay = act_layout(d)
    c, t = _company(act), _tenant(act)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.top_margin = sec.bottom_margin = Mm(15)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(2)

    def para(text="", bold=False, align=None, size=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        return p

    def header(lines):
        para(lines[0], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
        for x in lines[1:]:
            para(x, align=WD_ALIGN_PARAGRAPH.CENTER)
        hdr = doc.add_table(rows=1, cols=2)
        hdr.cell(0, 0).text = d.get("city") or ""
        hdr.cell(0, 1).text = _act_date(act)
        hdr.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        para()

    def table(head, rows, widths_mm, num_from):
        widths = [Mm(w) for w in widths_mm]
        tbl = doc.add_table(rows=1, cols=len(head))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w.twips)))
        for i, h in enumerate(head):
            cell = tbl.rows[0].cells[i]
            cell.text = ""
            cell.paragraphs[0].add_run(h).bold = True
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "EEEEEE")
            cell._tc.get_or_add_tcPr().append(shd)
        for r in rows:
            cells = tbl.add_row().cells
            for j, v in enumerate(r):
                cells[j].text = str(v)
                if j >= num_from:
                    cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in tbl.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = w
                for p in row.cells[j].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)

    def signatures(full=True):
        para()
        sig = doc.add_table(rows=2, cols=2)
        for col, (label, party) in enumerate([("Від Орендодавця:", c), ("Від Орендаря:", t)]):
            cell = sig.cell(0, col)
            cell.text = ""
            cell.paragraphs[0].add_run(label).bold = True
            if full:
                for x in _party(party):
                    cell.add_paragraph(x).runs[0].font.size = Pt(9)
            sc = sig.cell(1, col)
            sc.text = ""
            sc.paragraphs[0].paragraph_format.space_before = Pt(24)
            sc.paragraphs[0].add_run(_sign_line(party))
            sc.add_paragraph("М.П.")

    # ---- аркуш 1: акт ----
    header([f"АКТ № {act['number']}", "приймання-передачі наданих послуг"])
    para(_intro(act)).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    main_rows = [[str(i), l["name"], l["unit"], fmt_num(l["qty"]), fmt_money(l["price"]), fmt_money(l["sum"])]
                 for i, l in enumerate(lay["main"], 1)]
    table(MAIN_HEAD, main_rows, [8, 95, 17, 18, 21, 21], 3)
    for a, b in _totals_rows(act):
        para(f"{a} {b}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(_words(act))
    if lay["util"]:
        para("Розрахунок суми компенсації комунальних послуг наведено в Деталізації (додаток до цього акта).")
    para(f"Форма оплати: {d['payment_label']}. Сторони претензій одна до одної не мають.")
    signatures()

    # ---- аркуш 2: деталізація ----
    if lay["util"]:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        header(["ДЕТАЛІЗАЦІЯ КОМУНАЛЬНИХ ПОСЛУГ", f"за {_ulabel(act)} — додаток до акта № {act['number']} від {_act_date(act)}"])
        para(_detail_intro(act)).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        table(DETAIL_HEAD, _detail_rows(lay), [8, 45, 30, 27, 15, 17, 19, 19], 4)
        for a, b in _detail_totals_rows(act, lay):
            para(f"{a} {b}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
        para(f"Сума компенсації комунальних послуг: {amount_in_words(lay['util_total'])}"
             + (f", у т.ч. ПДВ {fmt_money(lay['util_vat'])} грн." if lay["util_vat"] else ", без ПДВ."))
        para("Ціни вказано без ПДВ. Кількість за лічильниками = (поточні − попередні показники) × коефіцієнт.",
             size=9)
        signatures(full=False)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
