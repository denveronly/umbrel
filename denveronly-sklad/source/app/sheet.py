"""Таблиця орендарів «як в Excel»: рядок на кожен склад, групи за підкатегоріями (Ангари, Навіси…),
підсумки «Безготівка» / «Готівка» по групі та загалом; керування підкатегоріями; вивантаження в .xlsx."""
import io
from datetime import date

from flask import Blueprint, flash, redirect, request, send_file, url_for

import calc
from audit import log
from db import categories, get_db, settings
from i18n import _

bp = Blueprint("sheet", __name__)
NO_CAT = None


def _contract(t):
    parts = []
    if t["contract_no"]:
        parts.append(f"№ {t['contract_no']}")
    if t["contract_date"]:
        parts.append(_("от") + f" {t['contract_date']}")
    return " ".join(parts)


def build(q=""):
    """Групи таблиці: [{"cat": row|None, "name", "rows": [...], "tot": {...}}], загальні підсумки."""
    db = get_db()
    s = settings()
    cats = categories()
    groups = {c["id"]: {"cat": c, "name": c["name"], "rows": []} for c in cats}
    groups[NO_CAT] = {"cat": None, "name": _("Без подкатегории"), "rows": []}

    tenants = {t["id"]: t for t in db.execute(
        "SELECT t.*, c.name AS company_name FROM tenants t LEFT JOIN companies c ON c.id=t.company_id "
        "WHERE t.active=1 ORDER BY t.name COLLATE NOCASE").fetchall()}
    contacts = {}
    for c in db.execute("SELECT * FROM tenant_contacts ORDER BY id"):
        contacts.setdefault(c["tenant_id"], []).append(c)
    whs = db.execute("SELECT * FROM warehouses ORDER BY name COLLATE NOCASE").fetchall()
    by_tenant = {}
    for w in whs:
        if w["tenant_id"] in tenants:
            by_tenant.setdefault(w["tenant_id"], []).append(w)

    def row(t, w):
        pt = t["payment_type"] if t else None
        price = calc.rent_price(w, pt, s) if (w and t) else None
        cost = price["total"] if price else 0
        cts = contacts.get(t["id"], []) if t else []
        hay = " ".join(str(x or "") for x in (
            [t["name"], t["edrpou"], t["contract_no"], t["company_name"], t["note"], t["contact"]] if t else [])
            + ([w["name"], w["note"]] if w else []) + [f"{c['name']} {c['phone']}" for c in cts]).lower()
        return {"t": t, "w": w, "contacts": cts, "pt": pt, "price": price,
                "bank": cost if pt in ("bank", "bank_vat") else 0, "cash": cost if pt == "cash" else 0,
                "area": (w["area"] or 0) if w else 0, "adj": (w["adjacent_area"] or 0) if w else 0, "hay": hay}

    for tid, t in tenants.items():
        tw = by_tenant.get(tid) or [None]
        for w in tw:
            groups[w["category_id"] if w and w["category_id"] in groups else NO_CAT]["rows"].append(row(t, w))
    free = [w for w in whs if w["tenant_id"] not in tenants]
    for w in free:
        groups[w["category_id"] if w["category_id"] in groups else NO_CAT]["rows"].append(row(None, w))

    ql = (q or "").strip().lower()
    out = []
    for g in groups.values():
        rows = [r for r in g["rows"] if not ql or ql in r["hay"] or ql in g["name"].lower()]
        # орендарі за назвою, їхні склади поруч; вільні склади — в кінці групи
        rows.sort(key=lambda r: (r["t"] is None, (r["t"]["name"].lower() if r["t"] else ""),
                                 (r["w"]["name"].lower() if r["w"] else "")))
        if not rows and (g["cat"] is None or ql):
            continue
        g["rows"] = rows
        g["tot"] = {k: calc.r2(sum(r[k] for r in rows)) for k in ("area", "adj", "bank", "cash")}
        g["rented"] = calc.r2(sum(r["area"] for r in rows if r["t"]))
        out.append(g)

    # нумерація як у таблиці: 1, 2.1, 2.2, 3… (номер орендаря + номер його складу, якщо складів кілька)
    nwh = {}
    for g in out:
        for r in g["rows"]:
            if r["t"]:
                nwh[r["t"]["id"]] = nwh.get(r["t"]["id"], 0) + 1
    num, seen, sub = 0, {}, {}
    for g in out:
        prev = None
        for r in g["rows"]:
            t = r["t"]
            if t:
                if t["id"] not in seen:
                    num += 1
                    seen[t["id"]] = num
                sub[t["id"]] = sub.get(t["id"], 0) + 1
                r["no"] = f"{seen[t['id']]}.{sub[t['id']]}" if nwh[t["id"]] > 1 else str(seen[t["id"]])
                r["first"] = not prev or not prev["t"] or prev["t"]["id"] != t["id"]
            else:
                r["no"] = ""
                r["first"] = True
            prev = r
    grand = {k: calc.r2(sum(g["tot"][k] for g in out)) for k in ("area", "adj", "bank", "cash")}
    grand["rented"] = calc.r2(sum(g["rented"] for g in out))
    return out, grand


# ---------------- підкатегорії ----------------

@bp.route("/categories", methods=["POST"])
def manage():
    db = get_db()
    a = request.form.get("action")
    cid = request.form.get("cid", type=int)
    name = request.form.get("name", "").strip()
    if a == "add" and name:
        db.execute("INSERT INTO warehouse_categories(name, sort) VALUES (?,?)",
                   (name, request.form.get("sort", type=int) or 100))
        log("Подкатегория добавлена", name, commit=False)
        flash(_("Подкатегория добавлена") + f": {name}")
    elif a == "save" and cid:
        if request.form.get("delete"):
            old = db.execute("SELECT name FROM warehouse_categories WHERE id=?", (cid,)).fetchone()
            db.execute("UPDATE warehouses SET category_id=NULL WHERE category_id=?", (cid,))
            db.execute("DELETE FROM warehouse_categories WHERE id=?", (cid,))
            log("Подкатегория удалена", old["name"] if old else cid, commit=False)
            flash(_("Подкатегория удалена, её склады — без подкатегории"))
        elif name:
            db.execute("UPDATE warehouse_categories SET name=?, sort=? WHERE id=?",
                       (name, request.form.get("sort", type=int) or 100, cid))
            log("Подкатегория изменена", name, commit=False)
            flash(_("Сохранено"))
    elif a == "assign":
        n = 0
        valid = {c["id"] for c in categories()}
        for k, v in request.form.items():
            if k.startswith("wc_"):
                new = int(v) if v and int(v) in valid else None
                cur = db.execute("SELECT category_id FROM warehouses WHERE id=?", (int(k[3:]),)).fetchone()
                if cur and cur["category_id"] != new:
                    db.execute("UPDATE warehouses SET category_id=? WHERE id=?", (new, int(k[3:])))
                    n += 1
        if n:
            log("Склады распределены по подкатегориям", str(n), commit=False)
        flash(_("Сохранено") + f": {n}")
    db.commit()
    return redirect(url_for("tenants", cats=1) + "#cats")


# ---------------- вивантаження в Excel ----------------

@bp.route("/tenants.xlsx")
def xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    groups, grand = build(request.args.get("q", ""))
    wb = Workbook()
    ws = wb.active
    ws.title = _("Арендаторы")[:31]
    head = ["№", _("Арендатор"), _("Объект"), _("S, м²"), _("S прилег., м²"), _("Примечание"), _("Договор"),
            _("до даты"), _("Контактное лицо"), _("Стоимость 1 (безнал)"), _("Стоимость 2 (наличка)"), _("Компания")]
    widths = [6, 30, 30, 10, 11, 24, 18, 14, 28, 15, 15, 20]
    thin, thick = Side(style="thin", color="999999"), Side(style="medium", color="000000")
    blue = PatternFill("solid", fgColor="C9DAF8")
    wrap = Alignment(wrap_text=True, vertical="bottom")
    money = '#,##0.00'
    ws.append(head)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        c.border = Border(bottom=thick)
    ws.freeze_panes = "A2"
    for g in groups:
        ws.append([None, g["name"], None, g["tot"]["area"], g["tot"]["adj"] or None, None, None, None, None,
                   g["tot"]["bank"], g["tot"]["cash"], None])
        r = ws.max_row
        for i in range(1, len(head) + 1):
            c = ws.cell(row=r, column=i)
            c.fill = blue
            c.font = Font(bold=True, size=13 if i in (2, 10, 11) else 11)
            c.border = Border(top=thick, bottom=thick)
        for i in (10, 11):
            ws.cell(row=r, column=i).number_format = money
        for x in g["rows"]:
            t, w = x["t"], x["w"]
            cts = "\n".join(f"{c['name'] or ''} {c['phone'] or ''}".strip() for c in x["contacts"])
            obj = (w["name"] + (f" — {w['note']}" if w["note"] else "")) if w else ""
            ws.append([x["no"], t["name"] if t else _("свободен"), obj, x["area"] or None, x["adj"] or None,
                       (t["note"] if t else "") or "", _contract(t) if t else "", (t["contract_end"] if t else "") or "",
                       cts, x["bank"] if t else None, x["cash"] if t else None, (t["company_name"] if t else "") or ""])
            r = ws.max_row
            for i in range(1, len(head) + 1):
                c = ws.cell(row=r, column=i)
                c.alignment = wrap
                c.border = Border(top=thick if x["first"] else thin, bottom=thin, left=thin, right=thin)
                if not t:
                    c.font = Font(color="888888", italic=True)
            for i in (10, 11):
                ws.cell(row=r, column=i).number_format = money
    ws.append([None, _("Итого"), None, grand["area"], grand["adj"] or None, None, None, None, None,
               grand["bank"], grand["cash"], None])
    r = ws.max_row
    for i in range(1, len(head) + 1):
        c = ws.cell(row=r, column=i)
        c.font = Font(bold=True)
        c.border = Border(top=thick)
    for i in (10, 11):
        ws.cell(row=r, column=i).number_format = money
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Orendari_{date.today():%Y-%m-%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def init_app(app):
    app.register_blueprint(bp)
